"""Geração de relatórios em CSV, JSON, Excel e PDF.

Os relatórios são a forma recomendada de revisar as recomendações **antes** de
qualquer operação de arquivo: exporte, confira com calma e só depois decida.
"""

from __future__ import annotations

import csv
import json
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from ..db.repository import Repository
from .fileops import format_bytes
from .models import BadFlag, Category

log = logging.getLogger(__name__)

FORMATS = {
    "csv": "Valores separados (CSV)",
    "json": "JSON",
    "xlsx": "Planilha do Excel",
    "pdf": "Documento PDF",
}


@dataclass
class ReportData:
    """Conteúdo do relatório, independente do formato de saída."""

    generated_at: str = field(default_factory=lambda: datetime.now().strftime("%d/%m/%Y %H:%M"))
    summary: dict = field(default_factory=dict)
    groups: list = field(default_factory=list)
    bad_photos: list = field(default_factory=list)
    folders: list = field(default_factory=list)
    settings_snapshot: dict = field(default_factory=dict)

    # ------------------------------------------------------------------ texto
    def summary_lines(self) -> list[tuple[str, str]]:
        s = self.summary
        return [
            ("Fotos analisadas", f"{s.get('photos', 0):n}"),
            ("Duplicatas exatas", f"{s.get('exact_duplicates', 0):n} arquivos em {s.get('exact_groups', 0):n} grupos"),
            ("Duplicatas visuais", f"{s.get('visual_duplicates', 0):n} arquivos em {s.get('visual_groups', 0):n} grupos"),
            ("Grupos de fotos muito semelhantes", f"{s.get('very_similar_groups', 0):n}"),
            ("Grupos de fotos semelhantes", f"{s.get('similar_groups', 0):n}"),
            ("Espaço potencialmente liberável", format_bytes(s.get("reclaimable_bytes", 0))),
            ("Espaço ocupado por cópias", format_bytes(s.get("redundant_bytes", 0))),
            ("Fotos com possível problema", f"{s.get('bad_photos', 0):n}"),
            ("Fotos sem EXIF", f"{s.get('no_exif', 0):n}"),
            ("Arquivos com erro de leitura", f"{s.get('errors', 0):n}"),
        ]

    def group_rows(self) -> list[dict]:
        rows: list[dict] = []
        for group in self.groups:
            for member in group.members:
                sig = member.signature
                rows.append(
                    {
                        "grupo": group.group_id,
                        "categoria": group.category.label,
                        "arquivo": sig.path,
                        "nome": sig.name,
                        "pasta": sig.folder,
                        "largura": sig.width,
                        "altura": sig.height,
                        "megapixels": round(sig.megapixels, 2),
                        "tamanho_bytes": sig.size,
                        "tamanho": format_bytes(sig.size),
                        "formato": sig.format,
                        "qualidade": round(sig.quality, 1),
                        "similaridade": round(member.similarity, 2),
                        "melhor_do_grupo": "sim" if member.is_reference else "não",
                        "recomendacao": "manter" if member.effective_choice == "keep" else "remover",
                        "motivo": member.reason,
                        "data_foto": sig.taken_at,
                        "camera": sig.camera,
                    }
                )
        return rows

    def bad_rows(self) -> list[dict]:
        rows = []
        for photo in self.bad_photos:
            flags = [f for f in (photo.get("bad_flags") or "").split(",") if f]
            labels = []
            for flag in flags:
                try:
                    labels.append(BadFlag(flag).label)
                except ValueError:
                    labels.append(flag)
            rows.append(
                {
                    "arquivo": photo["path"],
                    "problemas": ", ".join(labels),
                    "largura": photo["width"],
                    "altura": photo["height"],
                    "tamanho": format_bytes(photo["size"]),
                    "qualidade": round(float(photo["quality"] or 0), 1),
                    "formato": photo["format"],
                }
            )
        return rows


def collect(repo: Repository, settings_snapshot: dict | None = None, max_groups: int = 5000) -> ReportData:
    """Reúne os dados do relatório a partir do banco."""
    data = ReportData(
        summary=repo.summary(),
        groups=repo.load_groups(limit=max_groups, order="category"),
        bad_photos=repo.filtered_photos({"only_bad": True}, order="quality_asc", limit=5000),
        folders=[f["path"] for f in repo.list_folders()],
        settings_snapshot=settings_snapshot or {},
    )
    return data


# ----------------------------------------------------------------- exportação
def export(data: ReportData, path: str | Path, fmt: str | None = None) -> Path:
    """Exporta o relatório no formato indicado (ou deduzido pela extensão)."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    fmt = (fmt or target.suffix.lstrip(".")).lower()
    if fmt == "csv":
        return _export_csv(data, target)
    if fmt == "json":
        return _export_json(data, target)
    if fmt in ("xlsx", "xls"):
        return _export_xlsx(data, target)
    if fmt == "pdf":
        return _export_pdf(data, target)
    raise ValueError(f"Formato de relatório não suportado: {fmt}")


def _export_csv(data: ReportData, target: Path) -> Path:
    rows = data.group_rows()
    with open(target, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow([f"Relatório PhotoDedupe - {data.generated_at}"])
        for label, value in data.summary_lines():
            writer.writerow([label, value])
        writer.writerow([])
        if rows:
            dict_writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
            dict_writer.writeheader()
            dict_writer.writerows(rows)
    return target


def _export_json(data: ReportData, target: Path) -> Path:
    payload = {
        "gerado_em": data.generated_at,
        "aplicativo": "PhotoDedupe",
        "pastas": data.folders,
        "resumo": data.summary,
        "configuracao": data.settings_snapshot,
        "grupos": [
            {
                "id": g.group_id,
                "categoria": g.category.value,
                "categoria_rotulo": g.category.label,
                "quantidade": g.size,
                "espaco_liberavel_bytes": g.wasted_bytes,
                "fotos": [
                    {
                        "arquivo": m.signature.path,
                        "resolucao": f"{m.signature.width}x{m.signature.height}",
                        "tamanho_bytes": m.signature.size,
                        "formato": m.signature.format,
                        "qualidade": round(m.signature.quality, 1),
                        "similaridade": round(m.similarity, 2),
                        "melhor_do_grupo": m.is_reference,
                        "recomendacao": m.effective_choice,
                        "motivo": m.reason,
                        "detalhes_tecnicos": m.detail,
                    }
                    for m in g.members
                ],
            }
            for g in data.groups
        ],
        "fotos_com_problema": data.bad_rows(),
    }
    target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return target


def _export_xlsx(data: ReportData, target: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "A exportação para Excel exige o pacote 'openpyxl'. Instale-o ou escolha CSV/JSON/PDF."
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Relatório PhotoDedupe"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = f"Gerado em {data.generated_at}"
    row = 4
    for label, value in data.summary_lines():
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 42

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")

    def write_sheet(name: str, rows: list[dict]) -> None:
        sheet = wb.create_sheet(name)
        if not rows:
            sheet["A1"] = "Nada a listar."
            return
        headers = list(rows[0].keys())
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header.replace("_", " ").capitalize())
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for r, item in enumerate(rows, start=2):
            for c, header in enumerate(headers, start=1):
                sheet.cell(row=r, column=c, value=item[header])
        for col, header in enumerate(headers, start=1):
            width = max(12, min(70, max(len(str(header)), *(len(str(it[header])) for it in rows[:200])) + 2))
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    write_sheet("Grupos", data.group_rows())
    write_sheet("Fotos com problema", data.bad_rows())
    wb.save(target)
    return target


def _export_pdf(data: ReportData, target: Path) -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "A exportação para PDF exige o pacote 'reportlab'. Instale-o ou escolha CSV/JSON/Excel."
        ) from exc

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("titulo", parent=styles["Title"], fontSize=20, spaceAfter=6)
    small = ParagraphStyle("pequeno", parent=styles["Normal"], fontSize=8, leading=10)

    doc = SimpleDocTemplate(
        str(target), pagesize=A4, title="Relatório PhotoDedupe",
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
    )
    story = [
        Paragraph("Relatório PhotoDedupe", title_style),
        Paragraph(f"Gerado em {data.generated_at}", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("Este relatório é apenas informativo. Nenhum arquivo foi alterado para gerá-lo.", small),
        Spacer(1, 10),
        Paragraph("Resumo da análise", styles["Heading2"]),
    ]

    summary_table = Table([[k, v] for k, v in data.summary_lines()], colWidths=[85 * mm, 80 * mm])
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#1F3A5F")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([summary_table, Spacer(1, 14), Paragraph("Grupos encontrados", styles["Heading2"])])

    if data.folders:
        story.insert(4, Paragraph("Pastas analisadas: " + "; ".join(data.folders), small))

    max_groups = 400
    for group in data.groups[:max_groups]:
        head = (
            f"{group.category.emoji} Grupo {group.group_id:03d} — {group.category.label} — "
            f"{group.size} fotos — liberável: {format_bytes(group.wasted_bytes)}"
        )
        story.append(Paragraph(head, styles["Heading4"]))
        rows = [["Arquivo", "Resolução", "Tamanho", "Qualidade", "Semelhança", "Recomendação"]]
        for m in group.members:
            sig = m.signature
            rows.append(
                [
                    Paragraph(sig.name + ("  ⭐" if m.is_reference else ""), small),
                    f"{sig.width}×{sig.height}",
                    format_bytes(sig.size),
                    f"{sig.quality:.0f}/100",
                    f"{m.similarity:.1f}%",
                    "MANTER" if m.effective_choice == "keep" else "remover",
                ]
            )
        table = Table(rows, colWidths=[62 * mm, 25 * mm, 22 * mm, 22 * mm, 22 * mm, 27 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ]
            )
        )
        story.extend([table, Spacer(1, 8)])

    if len(data.groups) > max_groups:
        story.append(
            Paragraph(
                f"... e mais {len(data.groups) - max_groups} grupos. Use a exportação em CSV ou Excel para a lista completa.",
                small,
            )
        )

    bad_rows = data.bad_rows()
    if bad_rows:
        story.extend([Spacer(1, 10), Paragraph("Fotos com possível problema", styles["Heading2"])])
        rows = [["Arquivo", "Problemas", "Resolução", "Qualidade"]]
        for item in bad_rows[:300]:
            rows.append([Paragraph(item["arquivo"], small), item["problemas"], f"{item['largura']}×{item['altura']}", f"{item['qualidade']:.0f}"])
        table = Table(rows, colWidths=[85 * mm, 45 * mm, 25 * mm, 20 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7A3E00")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ]
            )
        )
        story.append(table)

    story.extend(
        [
            Spacer(1, 12),
            Paragraph(
                "Privacidade: todas as fotos foram analisadas localmente neste computador. "
                "Nenhuma imagem foi enviada para a internet.",
                small,
            ),
        ]
    )
    doc.build(story)
    return target
